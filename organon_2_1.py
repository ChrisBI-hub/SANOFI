import pandas as pd
import sqlalchemy as sa
from sqlalchemy import create_engine
import warnings
import openpyxl as pxl
from openpyxl.styles import Font, Alignment, PatternFill

warnings.filterwarnings('ignore')

# --- CONFIGURACIÓN ---
SERVER = "150.1.1.152"
DATABASE = "SIR"
USER = "ConsultaBD"
PASS = "5D$bc#kM&5W2T8J40?s%"
MES_VALIDACION = 2
ANIO_VALIDACION = 2026
FECHA_INICIO_VALIDACION = None
FECHA_FIN_VALIDACION = None
ETIQUETA_PERIODO = None

def obtener_engine():
    conn_url = (
        f"mssql+pyodbc://{USER}:{PASS}@{SERVER}/{DATABASE}?"
        f"driver=ODBC+Driver+18+for+SQL+Server&"
        f"Encrypt=yes&TrustServerCertificate=yes&"
        f"login_timeout=60"
    )
    return create_engine(conn_url)

fecha_ref_expr = (
    "COALESCE("
    "TRY_CONVERT(DATE, REF.CuentaG_FechaFactura, 103), "
    "TRY_CONVERT(DATE, REF.CuentaG_FechaFactura, 23), "
    "TRY_CONVERT(DATE, REF.CuentaG_FechaFactura, 120)"
    ")"
)
fecha_sabana_expr = (
    "COALESCE("
    "TRY_CONVERT(DATE, SABANA.CuentaG_FechaFactura, 103), "
    "TRY_CONVERT(DATE, SABANA.CuentaG_FechaFactura, 23), "
    "TRY_CONVERT(DATE, SABANA.CuentaG_FechaFactura, 120)"
    ")"
)

if FECHA_INICIO_VALIDACION and FECHA_FIN_VALIDACION:
    filtro_fecha_ref = (
        f"{fecha_ref_expr} >= CONVERT(DATE, '{FECHA_INICIO_VALIDACION.replace('-', '')}', 112) "
        f"AND {fecha_ref_expr} <= CONVERT(DATE, '{FECHA_FIN_VALIDACION.replace('-', '')}', 112)"
    )
    filtro_fecha_sabana = (
        f"{fecha_sabana_expr} >= CONVERT(DATE, '{FECHA_INICIO_VALIDACION.replace('-', '')}', 112) "
        f"AND {fecha_sabana_expr} <= CONVERT(DATE, '{FECHA_FIN_VALIDACION.replace('-', '')}', 112)"
    )
    etiqueta_salida = ETIQUETA_PERIODO or f"{FECHA_INICIO_VALIDACION}_{FECHA_FIN_VALIDACION}"
else:
    filtro_fecha_ref = (
        f"YEAR({fecha_ref_expr}) = {ANIO_VALIDACION} "
        f"AND MONTH({fecha_ref_expr}) = {MES_VALIDACION}"
    )
    filtro_fecha_sabana = (
        f"YEAR({fecha_sabana_expr}) = {ANIO_VALIDACION} "
        f"AND MONTH({fecha_sabana_expr}) = {MES_VALIDACION}"
    )
    etiqueta_salida = f"{ANIO_VALIDACION}_{MES_VALIDACION:02d}"

query1 = f"""
SELECT
    [Tipo Operación Desc] as [Tipo de Operación],
    Importador,
    COALESCE(
        TRY_CONVERT(DATETIME, [Pedimento Fecha Pago], 103),
        TRY_CONVERT(DATETIME, [Pedimento Fecha Pago], 23),
        TRY_CONVERT(DATETIME, [Pedimento Fecha Pago], 120)
    ) as [Fecha de Pago],
    [Aduana/Sección Despacho],
    REF.Referencia,
    Pedimento,
    Mercancía,
    [Clave Pedimento],
    COALESCE(
        TRY_CONVERT(DATETIME, FechaDetalle, 103),
        TRY_CONVERT(DATETIME, FechaDetalle, 23),
        TRY_CONVERT(DATETIME, FechaDetalle, 120)
    ) AS [Fecha de factura],
    NoFactura AS [No Factura],
    [NomProveedor] AS [Proveedor],
    Concepto AS [Descripcion del Gasto],
    DocUUID AS [UUID de la Factura],
    Subtotal AS [Subtotal],
    IVA AS [IVA],
    RetFlete AS [Retension],
    ImporteME AS [Total]
FROM Admin.Admin_VT_Pagos_Hechos AS PAGOS
JOIN [Admin].[SIR_VT_Sabana_Pedimento_ABC] AS REF 
    ON CONCAT(SUBSTRING(TRIM(REF.[CuentaG_Folio_Num_Factura]),1,1),TRIM(REF.[CuentaG_FolioFactura])) = PAGOS.[CTA. Gastos]
WHERE (
        REF.Cliente IN ('SANOFI PASTEUR, S.A DE C.V.', 'AZTECA VACUNAS, SA DE CV')
        OR (REF.Cliente LIKE '%AVENTIS%' AND REF.[EJE UNIDAD DE NEGOCIO] LIKE 'GENMED%')
    )
    AND {filtro_fecha_ref}
    AND REF.[Sucursal] LIKE '%CORRESPONSALIAS%'
    AND NomProveedor NOT LIKE '%TESORERIA DE LA FEDERACION%' 
    AND Concepto NOT LIKE '%RECTIFICACION DE PEDIMENTO%'
    AND Concepto NOT LIKE '%SERVICIOS CORRESPONSALES%'
"""

query2 = f"""
SELECT
    [Tipo Operación Desc] as [Tipo de Operación],
    Importador,
    COALESCE(
        TRY_CONVERT(DATETIME, SABANA.[dFechaPago], 103),
        TRY_CONVERT(DATETIME, SABANA.[dFechaPago], 23),
        TRY_CONVERT(DATETIME, SABANA.[dFechaPago], 120)
    ) as [Fecha de Pago],
    SABANA.[Aduana/Sección Despacho] as [Aduana/Sección Despacho],
    SABANA.Referencia,
    Pedimento as [Pedimento],
    SABANA.[Mercancía] as [Mercancía],
    SABANA.[Clave Pedimento] as [Clave Pedimento],
    COALESCE(
        TRY_CONVERT(DATETIME, SABANA.CuentaG_FechaFactura, 103),
        TRY_CONVERT(DATETIME, SABANA.CuentaG_FechaFactura, 23),
        TRY_CONVERT(DATETIME, SABANA.CuentaG_FechaFactura, 120)
    ) AS [Fecha de factura],
    CONCAT(SUBSTRING(TRIM([CuentaG_Folio_Num_Factura]),1,1),TRIM([CuentaG_FolioFactura])) AS [No Factura],
    'ABSOLUTE BROKERAGE CUSTOMS, S.C.' AS [Proveedor],
    'HONORARIOS DE AGENCIA' AS [Descripcion del Gasto],
    TIMBRE.sUUID AS [UUID de la Factura],
    SubtotalCuentadeGastosServicioAA AS [Subtotal],
    IvaCuentadeGastos AS [IVA],
    0 AS [Retension],
    ISNULL(TotaldeGastos, SABANA.CuentaG_SaldoTotal) AS [Total]
FROM [Admin].SIR_VT_Sabana_Pedimento_ABC SABANA
LEFT JOIN [Admin].[ADMINO_15_CUENTAS_GASTOS] CUENTASG
    ON CONCAT(SUBSTRING(TRIM(SABANA.[CuentaG_Folio_Num_Factura]),1,1),TRIM(SABANA.[CuentaG_FolioFactura])) = CONCAT(CUENTASG.sPrefijo, CUENTASG.nNumero)
LEFT JOIN [Admin].[ADMINO_31_TIMBRE_FACTURA] TIMBRE
    ON CUENTASG.nIdCtaGastos15 = TIMBRE.nIdCtaGastos15
WHERE (
        SABANA.Cliente IN ('SANOFI PASTEUR, S.A DE C.V.', 'AZTECA VACUNAS, SA DE CV')
        OR (SABANA.Cliente LIKE '%AVENTIS%' AND SABANA.[EJE UNIDAD DE NEGOCIO] LIKE 'GENMED%')
    )
    AND {filtro_fecha_sabana}
    AND SABANA.[Sucursal] LIKE '%CORRESPONSALIAS%'
"""

engine = obtener_engine()
try:
    with engine.connect() as conn:
        print("📥 Extrayendo Gastos y Honorarios...")
        df1 = pd.read_sql(sa.text(query1), conn)
        df2 = pd.read_sql(sa.text(query2), conn)
except Exception as e:
    print(f"❌ Error: {e}")
    exit()

df_total = pd.concat([df1, df2], ignore_index=True)

# --- NUEVO FILTRO: ELIMINAR IMPUESTOS AL COMERCIO EXTERIOR PARA MNS Y LT ---
print("🧹 Limpiando entradas de impuestos para corresponsalías (MNS y LT)...")
# Convertimos a string y mayúsculas para asegurar que las comparaciones no fallen por tipeo
mask_ref_mns_lt = df_total['Referencia'].astype(str).str.upper().str.startswith(('MNS', 'LT'))
mask_gasto_impuestos = df_total['Descripcion del Gasto'].astype(str).str.upper().str.contains('IMPUESTOS AL COMERCIO EXTERIOR')

# Mantenemos solo las filas que NO (~) cumplan con ambas condiciones a la vez
df_total = df_total[~(mask_ref_mns_lt & mask_gasto_impuestos)]

# Procesamiento de fechas
for col in ['Fecha de Pago', 'Fecha de factura']:
    df_total[col] = pd.to_datetime(df_total[col], errors='coerce')
    df_total[col] = df_total[col].apply(lambda x: x.strftime('%d/%m/%Y') if pd.notnull(x) else '')

df_total = df_total.sort_values(by=['Referencia', 'Tipo de Operación'])

# --- EXCEL (SOLO CONGLOMERADO) ---
archivo_base = f"REPORTE_ADUANAL_SANOFI_{etiqueta_salida}.xlsx"

# Guardamos SOLO la hoja de resumen
df_total.to_excel(archivo_base, sheet_name='RESUMEN_TOTAL', index=False)

# Aplicar formato básico al archivo base
wb = pxl.load_workbook(archivo_base)
ws = wb.active
header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
font_white = Font(bold=True, color="FFFFFF")

for col in ws.columns:
    ws.column_dimensions[col[0].column_letter].width = 25
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = font_white
    cell.alignment = Alignment(horizontal="center")

wb.save(archivo_base)
print(f"✅ Reporte base generado exitosamente: {archivo_base}")
