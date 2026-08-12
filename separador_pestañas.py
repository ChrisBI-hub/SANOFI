import pandas as pd
import re
import openpyxl as pxl
from openpyxl.styles import Font, Alignment, PatternFill
import os

def separar_por_referencias(ruta_archivo):
    if not os.path.exists(ruta_archivo):
        print(f"❌ El archivo {ruta_archivo} no existe.")
        return

    print(f"📂 Abriendo archivo para separar: {ruta_archivo}")
    
    # --- MEJORA: Leer la primera hoja disponible independientemente del nombre ---
    try:
        # Primero intentamos con el nombre estándar
        df = pd.read_excel(ruta_archivo, sheet_name='RESUMEN_TOTAL')
    except ValueError:
        # Si falla, leemos la primera hoja (índice 0)
        print("⚠️ No se halló 'RESUMEN_TOTAL', leyendo la primera hoja del archivo...")
        df = pd.read_excel(ruta_archivo, sheet_name=0)
    
    # Crear un nuevo Excel Writer
    # 'replace' asegura que si la hoja ya existe, la sobreescriba
    with pd.ExcelWriter(ruta_archivo, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
        
        # Agrupar por referencia y crear hojas
        for ref, grupo in df.groupby('Referencia'):
            if pd.notnull(ref) and str(ref).strip() != "":
                # Limpiar nombre de la hoja
                nombre_hoja = re.sub(r'[\\/*?:\[\]]', '-', str(ref))[:31]
                grupo.to_excel(writer, sheet_name=nombre_hoja, index=False)
    
    # --- APLICAR FORMATO ---
    wb = pxl.load_workbook(ruta_archivo)
    header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
    font_white = Font(bold=True, color="FFFFFF")

    for ws in wb.worksheets:
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 25
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = font_white
            cell.alignment = Alignment(horizontal="center")

    wb.save(ruta_archivo)
    print(f"✨ ¡Listo! Archivo {ruta_archivo} separado por referencias.")

if __name__ == "__main__":
    separar_por_referencias("Reporte_SANOFI_07.xlsx")
