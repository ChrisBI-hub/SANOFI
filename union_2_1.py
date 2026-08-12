import pandas as pd
import numpy as np
import os
import glob
import warnings
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill

warnings.filterwarnings('ignore')

# --- CONFIGURACIÓN DE RUTAS ---
PATH_BASE = "/home/christian/Documentos/SANOFI"
ARCHIVO_ADUANAL = f"{PATH_BASE}/REPORTE_ADUANAL_SANOFI_2026_02.xlsx"
ARCHIVO_MAESTRO = f"{PATH_BASE}/REPORTE_MAESTRO_2026_02.xlsx"
ARCHIVO_SALIDA = f"{PATH_BASE}/Reporte_SANOFI.xlsx"

def limpiar_monto(serie):
    """Limpia strings de moneda y los convierte a float de forma segura"""
    return (serie.astype(str)
            .replace(r'[\$,\s]', '', regex=True)
            .replace(['nan', 'None', ''], '0')
            .astype(float)
            .round(2))

def unificar_reportes():
    # --- VALIDACIÓN DE EXISTENCIA DE ARCHIVOS DE CORRESPONSALÍAS ---
    patron_laredo = os.path.join(PATH_BASE, "ANALISIS_LT_*.xlsx")
    archivos_laredo = glob.glob(patron_laredo)
    
    existe_maestro = os.path.exists(ARCHIVO_MAESTRO)
    existe_laredo = len(archivos_laredo) > 0
    
    if not os.path.exists(ARCHIVO_ADUANAL):
        print(f"❌ Error: El archivo base {ARCHIVO_ADUANAL} no existe.")
        return

    print("🚀 Iniciando unificación de reportes...")
    if not existe_maestro and not existe_laredo:
        print("ℹ️ No hay operaciones en corresponsalías que agregar. Se generará el reporte con la base aduanal.")

    # 1. Leer Base Aduanal
    df_base = pd.read_excel(ARCHIVO_ADUANAL, sheet_name='RESUMEN_TOTAL')
    columnas_originales = df_base.columns.tolist()
    df_base['Ref_Key'] = df_base['Referencia'].astype(str).str.strip().str.upper()
    df_base['Total_Key'] = limpiar_monto(df_base['Total'])
    
    # 2. Procesar Maestro (Manzanillo) si existe
    if existe_maestro:
        print("✅ Procesando Archivo Maestro (Manzanillo)...")
        df_maestro = pd.read_excel(ARCHIVO_MAESTRO)
        if df_maestro.empty:
            print("ℹ️ Archivo Maestro vacío. Se omite el cruce de Manzanillo.")
            df_unificado = df_base.copy()
            for col in ['FP_Port', 'Desc_Port', 'Fact_Port', 'Prov_Port']:
                df_unificado[col] = np.nan
        else:
            df_maestro['Ref_Key'] = df_maestro['Referencia_Original'].astype(str).str.strip().str.upper()
            df_maestro['Total_Key'] = limpiar_monto(df_maestro['Importe'])
            
            df_portal_sub = df_maestro[['Ref_Key', 'Total_Key', 'Fecha de Pago', 'Descripcion_Consolidada', 'Factura', 'Emisor_fac']].copy()
            df_portal_sub.columns = ['Ref_Key', 'Total_Key', 'FP_Port', 'Desc_Port', 'Fact_Port', 'Prov_Port']
            df_unificado = pd.merge(df_base, df_portal_sub, on=['Ref_Key', 'Total_Key'], how='left')
    else:
        df_unificado = df_base.copy()
        for col in ['FP_Port', 'Desc_Port', 'Fact_Port', 'Prov_Port']:
            df_unificado[col] = np.nan

    # Aplicar datos de Manzanillo iniciales
    df_unificado['Fecha de Pago'] = df_unificado['FP_Port'].fillna(df_unificado['Fecha de Pago'])
    df_unificado['Mercancía'] = df_unificado['Desc_Port'].fillna(df_unificado['Mercancía'])
    df_unificado['No Factura'] = df_unificado['Fact_Port'].fillna(df_unificado['No Factura'])
    df_unificado['Proveedor'] = df_unificado['Prov_Port'].fillna(df_unificado['Proveedor'])

    # 3. Procesar Laredo si existe
    if existe_laredo:
        archivo_laredo = archivos_laredo[0]
        print(f"✅ Procesando Archivo Laredo: {os.path.basename(archivo_laredo)}")
        df_laredo = pd.read_excel(archivo_laredo)
        df_laredo.replace("S/D", np.nan, inplace=True)
        
        df_laredo['Ref_Key'] = df_laredo['Referencia'].astype(str).str.strip().str.upper()
        df_laredo_sub = df_laredo[['Ref_Key', 'Fecha Pago', 'Descripcion']].drop_duplicates(subset=['Ref_Key'])
        df_laredo_sub.columns = ['Ref_Key', 'FP_Laredo', 'Desc_Laredo']
        
        df_unificado = pd.merge(df_unificado, df_laredo_sub, on='Ref_Key', how='left')
        df_unificado['Fecha de Pago'] = df_unificado['Fecha de Pago'].fillna(df_unificado['FP_Laredo'])
        df_unificado['Mercancía'] = df_unificado['Mercancía'].fillna(df_unificado['Desc_Laredo'])

    # 4. Propagación por Pedimento (Crucial para cubrir todas las partidas)
    print("🔄 Propagando datos por número de Pedimento...")
    df_unificado['Ped_Key'] = df_unificado['Pedimento'].astype(str).str.strip()
    
    map_fechas = df_unificado.dropna(subset=['Fecha de Pago']).set_index('Ped_Key')['Fecha de Pago'].to_dict()
    map_mercancia = df_unificado.dropna(subset=['Mercancía']).set_index('Ped_Key')['Mercancía'].to_dict()

    df_unificado['Fecha de Pago'] = df_unificado.apply(lambda x: map_fechas.get(x['Ped_Key'], x['Fecha de Pago']), axis=1)
    df_unificado['Mercancía'] = df_unificado.apply(lambda x: map_mercancia.get(x['Ped_Key'], x['Mercancía']), axis=1)

    # 5. Estructura Final y Formato
    df_final = df_unificado[columnas_originales]
    df_final.to_excel(ARCHIVO_SALIDA, index=False)
    
    def aplicar_formato(ruta):
        wb = load_workbook(ruta)
        ws = wb.active
        header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
        font_white = Font(bold=True, color="FFFFFF")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 25
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = font_white
            cell.alignment = Alignment(horizontal="center", vertical="center")
        wb.save(ruta)

    aplicar_formato(ARCHIVO_SALIDA)
    
    # 6. Limpieza de temporales de Laredo
    if existe_laredo:
        try:
            os.remove(archivos_laredo[0])
            print(f"🗑️ Temporal de Laredo eliminado.")
        except: pass

    print(f"✅ Reporte unificado generado exitosamente: {ARCHIVO_SALIDA}")

if __name__ == "__main__":
    unificar_reportes()
